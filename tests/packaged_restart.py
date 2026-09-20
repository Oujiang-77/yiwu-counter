"""Verify persistence across a real frozen-program restart and inspect delivered files."""
import http.cookiejar
import io
import json
import subprocess
import time
import urllib.request
import zipfile
from pathlib import Path
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
data=root/'test-packaged'
process=subprocess.Popen([str(root/'dist'/'档口开单系统'/'档口开单系统.exe'),'--data-dir',str(data),'--port','18764','--no-browser'],creationflags=subprocess.CREATE_NO_WINDOW)
try:
    for attempt in range(40):
        try:
            state=json.loads((data/'runtime.json').read_text('utf8'))
            opener=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
            opener.open(state['url']+'/?session='+state['token'],timeout=1).read()
            break
        except (OSError,ValueError):time.sleep(.25)
    with opener.open(state['url']+'/api/state') as response:saved=json.load(response)
    assert len(saved['products'])==2
    assert len(saved['draft']['selected'])==2
    assert len(json.load(opener.open(state['url']+'/api/orders')))==2
    with zipfile.ZipFile(data/'two-factory.zip') as archive:
        assert len(archive.namelist())==2
        for entry in archive.namelist():
            book=load_workbook(io.BytesIO(archive.read(entry)),data_only=True)
            assert len(book.sheetnames)==1
            assert book.active['M4'].value==book.active['J4'].value*book.active['L4'].value
    opener.open(urllib.request.Request(state['url']+'/api/shutdown',method='POST',headers={'X-Counter-Request':'1'})).read()
    process.wait(timeout=15)
    assert process.returncode==0
    package=root.parent/'outputs'/'档口开单系统-0.1.0-Windows-x64.zip'
    with zipfile.ZipFile(package) as archive:
        assert archive.testzip() is None
        names=archive.namelist()
        assert not any(n.endswith(('.sqlite3','.log')) or '/test-' in n or 'runtime.json' in n for n in names)
        assert any(n.endswith('.onnx') for n in names)
    print('PASS: frozen restart retains products/draft/history; ZIP workbooks correct; release contains models and no business database')
finally:
    if process.poll() is None:process.terminate();process.wait(timeout=10)
