import io,json,sys,unittest,zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from test_local import LocalAppTests,product
from openpyxl import Workbook,load_workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image
from sheets import preview_import,template_bytes,FIELDS
from storage import validate_product,Store

class ProductMediaTests(unittest.TestCase):
    setUp=LocalAppTests.setUp
    tearDown=LocalAppTests.tearDown
    post=LocalAppTests.post
    def picture(self,color='red'):
        raw=io.BytesIO();Image.new('RGB',(25,25),color).save(raw,'PNG');return raw.getvalue()
    def upload(self,color='red'):
        r=self.client.post('/api/images',files={'file':('p.png',self.picture(color),'image/png')},headers=self.headers)
        self.assertEqual(r.status_code,200);return r.json()['image']
    def test_multi_images_save_restore_and_limits(self):
        main=[self.upload() for _ in range(5)];boxes=[self.upload('blue') for _ in range(5)]
        data=dict(product(),images=main,boxImages=boxes)
        r=self.post('/api/products',data);self.assertEqual(r.status_code,200,r.text)
        saved=self.client.get('/api/state').json()['products'][0]
        self.assertEqual(saved['images'],main);self.assertEqual(saved['boxImages'],boxes)
        self.assertTrue(all(k not in saved for k in ('area','inner','min')))
        bad=dict(data,images=main+[self.upload()]);self.assertEqual(self.post('/api/products',bad).status_code,400)
        self.assertEqual(self.post('/api/products',dict(data,boxImages=['f'*32+'.jpg'])).status_code,400)
        backup=self.post('/api/backups',{}).json()['name']
        r=self.post('/api/backups/restore',{'name':backup,'confirmation':'恢复备份'});self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(self.client.get('/api/state').json()['products'][0]['boxImages'],boxes)
        # Removing one primary photo preserves the other group.
        data['images']=main[1:]
        r=self.client.put('/api/products/'+str(saved['id']),json=data,headers=self.headers)
        self.assertEqual(r.json()['image'],main[1]);self.assertEqual(r.json()['boxImages'],boxes)

    def test_old_record_normalization(self):
        name=self.upload();legacy=dict(product(),image=name,min=999)
        store=self.app.state.store
        with store.connect() as c:
            c.execute('INSERT INTO products(code,factory,factory_code,data,created_at,updated_at) VALUES(?,?,?,?,?,?)',('A-13','测试电器厂','000138',json.dumps(legacy),'now','now'))
        p=Store(self.temp.name).products()[0]
        self.assertEqual(p['images'],[name]);self.assertEqual(p['boxImages'],[]);self.assertNotIn('min',p)
        self.assertEqual(store.order_items([{'id':p['id'],'quantity':1}])[0]['quantity'],1)

    def test_import_groups_and_mapping_changes(self):
        w=Workbook();s=w.active;s.append(['销售编码','厂家编码','商品名称','厂家名称','装箱数','出厂价','主图','彩盒'])
        s.append(['P-01','001','图片商品','图片厂',12,2])
        for col in ('G','H'):
            for _ in range(5):s.add_image(XLImage(io.BytesIO(self.picture())),col+'2')
        p=Path(self.temp.name)/'media.xlsx';w.save(p)
        cache={};r=preview_import(self.app.state.store,p,{'sheet':'Sheet'},cache)
        self.assertEqual(r['valid'],1);data=r['rows'][0]['data']
        self.assertEqual(len(data['images']),5);self.assertEqual(len(data['boxImages']),5)
        mapping=dict(r['mapping']);mapping['images'],mapping['boxImages']=7,6
        swapped=preview_import(self.app.state.store,p,{'sheet':'Sheet','mapping':mapping},cache)
        self.assertEqual(swapped['rows'][0]['data']['images'],data['boxImages'])
        template=load_workbook(io.BytesIO(template_bytes()))
        headers=[c.value for c in template.active[1]]
        self.assertIn('彩盒',headers);self.assertIn('主图',headers)
        self.assertFalse(any(k in headers for k in ('厂家地区','中包数','起订件数')))

    def test_custom_export_persistence_photos_and_totals(self):
        main=[self.upload() for _ in range(5)];boxes=[self.upload('blue') for _ in range(5)]
        saved=self.post('/api/products',dict(product(),images=main,boxImages=boxes)).json()
        columns=['amount','boxImages','name','images','quantity']
        r=self.client.put('/api/export-settings',json={'columns':columns},headers=self.headers);self.assertEqual(r.status_code,200)
        self.assertEqual(Store(self.temp.name).get_value('exportColumns'),columns)
        r=self.post('/api/orders',{'items':[{'id':saved['id'],'quantity':2}],'mode':'single'});self.assertEqual(r.status_code,200,r.text)
        b=load_workbook(io.BytesIO(self.client.get(r.json()['url']).content));s=b.active
        self.assertEqual([c.value for c in s[3]],['金额','彩盒','商品名称','主图','件数'])
        self.assertEqual(s['A4'].value,1500);self.assertEqual(s['A5'].value,'=SUM(A4:A4)');self.assertEqual(s['E5'].value,'=SUM(E4:E4)')
        self.assertEqual(len(s._images),10)
        self.assertEqual(sorted(im.anchor._from.col for im in s._images),[1]*5+[3]*5)
        for bad in ([],['name','name'],['area'],['inner'],['min']):
            r=self.client.put('/api/export-settings',json={'columns':bad},headers=self.headers);self.assertEqual(r.status_code,400)

if __name__=='__main__':unittest.main()
