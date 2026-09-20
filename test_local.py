import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent))
from fastapi.testclient import TestClient
from openpyxl import Workbook,load_workbook
from PIL import Image
from storage import Store
from app import make_app
from recognition import parse_lines

def product(code='A-13',factory='测试电器厂'):
    return dict(code=code,factoryCode='000138',name='测试风扇',factory=factory,area='浙江',size='10cm',battery='1200mAh',charger='Type-C',pack=60,inner=12,volume=.085,price=12.5,min=2)

class LocalAppTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.app=make_app(self.temp.name,'test-token')
        self.client=TestClient(self.app,base_url='http://127.0.0.1')
        self.client.get('/?session=test-token')
        self.headers={'X-Counter-Request':'1'}
    def tearDown(self):self.client.close();self.temp.cleanup()
    def post(self,path,data):return self.client.post(path,json=data,headers=self.headers)
    def test_persistence_and_duplicate(self):
        self.assertEqual(self.post('/api/products',product()).status_code,200)
        self.assertEqual(self.post('/api/products',product()).status_code,400)
        saved=Store(self.temp.name).products();self.assertEqual(len(saved),1);self.assertEqual(saved[0]['factoryCode'],'000138')
    def test_export_and_minimum(self):
        p=self.post('/api/products',product()).json()
        self.assertEqual(self.post('/api/orders',{'items':[{'id':p['id'],'quantity':1}]}).status_code,400)
        r=self.post('/api/orders',{'items':[{'id':p['id'],'quantity':3}],'mode':'single'});self.assertEqual(r.status_code,200,r.text)
        binary=self.client.get(r.json()['url']).content
        book=load_workbook(io.BytesIO(binary));ws=book.active
        self.assertEqual(ws['L4'].value,3);self.assertEqual(ws['M4'].value,180);self.assertEqual(ws['O4'].value,2250)
        # Changing a product does not alter an existing order snapshot.
        data=product();data['price']=99
        self.client.put('/api/products/'+str(p['id']),json=data,headers=self.headers)
        self.assertEqual(self.client.get('/api/orders').json()[0]['items'][0]['price'],12.5)
    def test_image_and_backup_restore(self):
        raw=io.BytesIO();Image.new('RGB',(20,20),'red').save(raw,'PNG')
        name=self.client.post('/api/images',files={'file':('p.png',raw.getvalue(),'image/png')},headers=self.headers).json()['image']
        data=product();data['image']=name;p=self.post('/api/products',data).json()
        self.client.put('/api/draft',json={'selected':[p['id']],'quantities':{str(p['id']):4}},headers=self.headers)
        backup=self.post('/api/backups',{}).json()['name']
        data=product('B-01','新厂家');self.post('/api/products',data)
        r=self.post('/api/backups/restore',{'name':backup,'confirmation':'恢复备份'});self.assertEqual(r.status_code,200,r.text)
        state=self.client.get('/api/state').json();self.assertEqual(len(state['products']),1);self.assertEqual(state['draft']['quantities'][str(p['id'])],4)
        self.assertEqual(self.client.get('/api/images/'+name).status_code,200)
    def test_import_mapping_and_invalid_rows(self):
        b=Workbook();s=b.active;s.title='厂家商品'
        s.append(['货号','品名','装箱数','单价']);s.append(['000001','第一款',48,18.8]);s.append(['000002','第二款','',9])
        raw=io.BytesIO();b.save(raw)
        upload=self.client.post('/api/import/upload',files={'file':('厂家.xlsx',raw.getvalue())},headers=self.headers).json()
        config=dict(batch=upload['batch'],sheet='厂家商品',header=1,factory='测试杯厂',prefix='B',start=1)
        r=self.post('/api/import/preview',config);self.assertEqual(r.status_code,200,r.text);preview=r.json();self.assertEqual(preview['valid'],1)
        r=self.post('/api/import/commit',dict(batch=upload['batch'],previewId=preview['previewId']));self.assertEqual(r.json()['count'],1)
        p=self.client.get('/api/state').json()['products'][0];self.assertEqual(p['code'],'B-01');self.assertEqual(p['factoryCode'],'000001')
        self.assertEqual(self.post('/api/import/commit',dict(batch=upload['batch'],previewId=preview['previewId'])).status_code,400)
    def test_atomic_import(self):
        s=self.app.state.store
        with self.assertRaises(ValueError):s.commit_import([product(),product()])
        self.assertEqual(s.products(),[])
    def test_import_embedded_image_and_leading_zero_number(self):
        from openpyxl.drawing.image import Image as XLImage
        from sheets import preview_import
        b=Workbook();s=b.active;s.title='资料'
        s.append(['销售编码','厂家编码','商品名称','厂家名称','装箱数','出厂价'])
        s.append(['X-01',123,'图片商品','图片厂',24,9.9]);s['B2'].number_format='000000'
        pic=io.BytesIO();Image.new('RGB',(40,40),'blue').save(pic,'PNG');pic.seek(0)
        s.add_image(XLImage(pic),'C2')
        path=Path(self.temp.name)/'import.xlsx';b.save(path)
        preview=preview_import(self.app.state.store,path,{'sheet':'资料','header':1})
        self.assertEqual(preview['valid'],1)
        row=preview['rows'][0]['data'];self.assertEqual(row['factoryCode'],'000123')
        self.assertTrue((Path(self.temp.name)/'images'/row['image']).is_file())
    def test_multi_factory_export_and_literal_text(self):
        first=product();first['name']='=1+1';first['price']=.1
        p1=self.post('/api/products',first).json();p2=self.post('/api/products',product('B-01','另一个厂家')).json()
        items=[{'id':p1['id'],'quantity':3},{'id':p2['id'],'quantity':2}]
        merged=self.post('/api/orders',{'items':items,'mode':'single'}).json()
        book=load_workbook(io.BytesIO(self.client.get(merged['url']).content));self.assertEqual(len(book.sheetnames),2)
        ws=book.worksheets[0];self.assertEqual(ws['A4'].value,'=1+1');self.assertEqual(ws['A4'].data_type,'s');self.assertEqual(ws['O4'].value,18)
        split=self.post('/api/orders',{'items':items,'mode':'multiple'}).json()
        with zipfile.ZipFile(io.BytesIO(self.client.get(split['url']).content)) as z:
            self.assertEqual(len(z.namelist()),2)
            for name in z.namelist():self.assertEqual(len(load_workbook(io.BytesIO(z.read(name))).sheetnames),1)
    def test_access_controls(self):
        self.assertEqual(self.client.post('/api/products',json=product()).status_code,403)
        self.assertEqual(self.client.post('/api/products',json=product(),headers={**self.headers,'Origin':'https://example.com'}).status_code,403)
        other=TestClient(self.app,base_url='http://127.0.0.1');self.assertEqual(other.get('/api/state').status_code,401);other.close()
        self.assertEqual(self.client.get('/api/state',headers={'Host':'malicious.example'}).status_code,403)
    def test_ocr_parsing_does_not_invent_unit(self):
        p=dict(product(),id=1)
        rows=parse_lines([{'text':'A-13 3','confidence':.95},{'text':'A-13 120个','confidence':.95}],[p])
        self.assertEqual(rows[0]['unit'],'');self.assertEqual(rows[0]['quantity'],3)
        self.assertEqual(rows[1]['quantity'],120);self.assertEqual(rows[1]['unit'],'个')
        self.assertFalse(rows[1]['reviewed'])

if __name__=='__main__':unittest.main(verbosity=2)
