from __future__ import annotations
import io
import re
import uuid
import zipfile
import posixpath
from xml.etree import ElementTree as ET
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from PIL import Image, ImageOps
from storage import validate_product

MAX_IMPORT_BYTES = 500 * 1024**2

def image_rows(path, sheet_name, store):
    """Read selected-sheet drawing images one at a time, without loading all media."""
    pics = {}; warnings = []
    ns = {'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
          'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
          'x':'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
          'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with zipfile.ZipFile(path) as z:
        def relationships(part):
            relpath = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part)+'.rels')
            if relpath not in z.namelist(): return {}
            return {e.attrib['Id']:posixpath.normpath(posixpath.join(posixpath.dirname(part),e.attrib['Target'])).lstrip('/')
                    for e in ET.fromstring(z.read(relpath)) if e.attrib.get('TargetMode') != 'External'}
        workbook = ET.fromstring(z.read('xl/workbook.xml'))
        sheet = next(e for e in workbook.findall('m:sheets/m:sheet',ns) if e.attrib['name']==sheet_name)
        part = relationships('xl/workbook.xml')[sheet.attrib['{'+ns['r']+'}id']]
        rels = relationships(part); drawings = []
        with z.open(part) as stream:
            for _, element in ET.iterparse(stream,events=('end',)):
                if element.tag == '{'+ns['m']+'}drawing': drawings.append(element.attrib['{'+ns['r']+'}id'])
                element.clear()
        for drawing in drawings:
            drawing_part = rels[drawing]; media = relationships(drawing_part)
            for anchor in ET.fromstring(z.read(drawing_part)):
                row = anchor.find('x:from/x:row',ns)
                blip = anchor.find('.//a:blip',ns)
                if blip is None: continue
                if row is None:
                    warnings.append('部分图片没有对应商品行，请导入后补充。'); continue
                number = int(row.text)+1
                try:
                    name = media[blip.attrib['{'+ns['r']+'}embed']]
                    if z.getinfo(name).file_size > 15*1024**2: raise ValueError('图片超过 15 MB')
                    col=int(anchor.find('x:from/x:col',ns).text)
                    bucket=pics.setdefault(number,[])
                    if sum(item['col']==col for item in bucket)>=5:
                        warnings.append(f'第 {number} 行同一图片列超过 5 张，仅保留前 5 张。');continue
                    bucket.append({'col':col,'image':save_image(z.read(name),store)})
                except Exception:
                    warnings.append(f'第 {number} 行图片无法导入（过大或格式不支持），请手动补充。')
    return pics,warnings

FIELDS={'code':'销售编码','factoryCode':'厂家编码','name':'商品名称','factory':'厂家名称','color':'颜色','size':'尺寸','cartonSize':'外箱尺寸','cartonWeight':'单箱重量（kg）','material':'材质','parameters':'产品参数','battery':'电池容量','charger':'充电头 / 接口','pack':'装箱数','volume':'单件体积','price':'出厂价','cat':'分类','notes':'备注','images':'主图','boxImages':'彩盒'}
ALIASES={'code':['销售编码','自定义编码','商家编码'],'factoryCode':['厂家编码','货号','型号','产品编码'],'name':['商品名称','产品名称','品名','名称'],'factory':['厂家名称','厂家','供应商'],'color':['颜色','产品颜色','色号'],'size':['尺寸','产品尺寸','规格'],'cartonSize':['外箱尺寸','箱子尺寸','外箱规格','箱规尺寸'],'cartonWeight':['单箱重量','单箱重量（kg）','单箱重量(kg)','毛重','每箱重量'],'material':['材质','产品材质'],'parameters':['产品参数','参数','规格参数'],'battery':['电池容量','电池'],'charger':['充电头 / 接口','充电头','接口'],'pack':['装箱数','装箱数量','每箱数量','箱规'],'volume':['单件体积','体积','每箱体积'],'price':['出厂价','单价','价格'],'cat':['分类','类别'],'notes':['备注'],'images':['主图','图片','商品图片','产品图片'],'boxImages':['彩盒','彩盒图片','包装图片']}

def clean_header(s):
    return re.sub(r'[\s\(（].*','',str(s).strip()).lower()

def save_image(data,store):
    if len(data)>15*1024**2:raise ValueError('图片超过 15 MB')
    with Image.open(io.BytesIO(data)) as im:
        if im.width*im.height>40_000_000:raise ValueError('图片分辨率过大')
        im=ImageOps.exif_transpose(im).convert('RGB')
        im.thumbnail((1800,1800))
        name=uuid.uuid4().hex+'.jpg'
        im.save(store.root/'images'/name,quality=88)
    return name

def inspect_book(path):
    if path.suffix=='.xls':
        import xlrd
        book=xlrd.open_workbook(path)
        return book.sheet_names()
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist())>2*1024**3:raise ValueError('表格展开后超过 2 GB，请拆分后导入')
        if sum(i.file_size for i in z.infolist() if not i.filename.startswith('xl/media/'))>200*1024**2:raise ValueError('表格文字数据展开后超过 200 MB，请拆分后导入')
    b=load_workbook(path,read_only=True,data_only=True)
    names=b.sheetnames;b.close();return names

def read_sheet(path,sheet,header_row,store):
    if not 1<=header_row<=50:raise ValueError('表头行须在第 1～50 行')
    pics={};warnings=[]
    if path.suffix=='.xls':
        import xlrd
        b=xlrd.open_workbook(path);s=b.sheet_by_name(sheet)
        if s.nrows>10000+header_row or s.ncols>100:raise ValueError('表格最多 10000 条数据、100 列，请拆分文件')
        vals=[[int(v) if isinstance(v,float) and v.is_integer() else v for v in s.row_values(i)] for i in range(s.nrows)]
        warnings.append('旧版 .xls 仅导入文字与数值，商品图片请另行上传。')
    else:
        b=load_workbook(path,data_only=True,read_only=True)
        s=b[sheet]
        vals=[]
        try:
            if (s.max_row or 0)>10000+header_row or (s.max_column or 0)>100:raise ValueError('表格最多 10000 条数据、100 列，请删除多余空行列或拆分文件')
            for row in s:
                if len(vals)>=10000+header_row or len(row)>100:raise ValueError('表格最多 10000 条数据、100 列，请拆分文件')
                line=[]
                for c in row:
                    v=c.value
                    if isinstance(v,(int,float)) and re.fullmatch('0{2,}',c.number_format or ''):
                        v=str(int(v)).zfill(len(c.number_format))
                    line.append(v)
                vals.append(line)
        finally:b.close()
        pics,warnings=image_rows(path,sheet,store)
    if len(vals)<header_row:raise ValueError('表头行超出工作表范围')
    heads=[str(v or '').strip() for v in vals[header_row-1]]
    return heads,[(i+1,vals[i]) for i in range(header_row,len(vals)) if any(v is not None and str(v).strip() for v in vals[i])],pics,warnings

def guess_mapping(headers):
    mapping={}
    for key,names in ALIASES.items():
        for i,h in enumerate(headers):
            if clean_header(h) in [clean_header(n) for n in names]:mapping[key]=i;break
    return mapping

def preview_import(store,path,config,cache=None):
    sheet=str(config['sheet']);header=int(config.get('header',1))
    key=(sheet,header)
    if cache is not None and key in cache: parsed=cache[key]
    else:
        parsed=read_sheet(path,sheet,header,store)
        if cache is not None:
            cache.clear();cache[key]=parsed
    heads,raw,pics,base_warnings=parsed
    warnings=list(base_warnings)
    mapping=config['mapping'] if 'mapping' in config else guess_mapping(heads)
    indices=[int(v) for v in mapping.values() if v is not None and str(v)!='']
    if len(indices)!=len(set(indices)):raise ValueError('同一列不能映射到多个字段')
    existing=store.products();codes={p['code'] for p in existing};factory_codes={(p['factory'],p['factoryCode'].upper()) for p in existing}
    result=[];prefix=str(config.get('prefix') or '').strip().upper()
    if prefix and not re.fullmatch(r'[A-Z0-9]{1,12}',prefix):raise ValueError('销售编码前缀仅支持 1～12 位字母或数字')
    try:start=int(config.get('start') or 1)
    except (ValueError,TypeError):raise ValueError('起始序号无效')
    if start<1:raise ValueError('起始序号须大于 0')
    for seq,(number,values) in enumerate(raw):
        data={}
        for key in FIELDS:
            if key in ('images','boxImages'):continue
            idx=mapping.get(key)
            if idx is not None and str(idx)!='':
                idx=int(idx)
                if idx<0 or idx>=len(heads):raise ValueError('列映射超出范围')
                val=values[idx] if idx<len(values) else None
                if isinstance(val,float) and val.is_integer():val=int(val)
                data[key]='' if val is None else val
        if not data.get('code') and prefix:data['code']=f'{prefix}-{start+seq:02d}'
        if not data.get('factory'):data['factory']=config.get('factory','')
        for k,v in {'volume':0,'cat':'日用百货'}.items():
            if data.get(k) in (None,''):data[k]=v
        images=pics.get(number,[])
        for key in ('images','boxImages'):
            col=mapping.get(key)
            if col is not None and str(col)!='':
                chosen=[entry['image'] for entry in images if entry['col']==int(col)]
            elif key=='images' and 'images' not in guess_mapping(heads):
                boxcol=mapping.get('boxImages')
                chosen=[entry['image'] for entry in images if boxcol is None or str(boxcol)=='' or entry['col']!=int(boxcol)]
            else:chosen=[]
            if len(chosen)>5:warnings.append(f'第 {number} 行{FIELDS[key]}超过 5 张，仅保留前 5 张。')
            data[key]=chosen[:5]
        data['image']=data['images'][0] if data['images'] else ''
        error=''
        try:
            data=validate_product(data)
            pair=(data['factory'],data['factoryCode'].upper())
            if data['code'] in codes:raise ValueError('销售编码与已有商品或本批次其他行重复')
            if pair in factory_codes:raise ValueError('同厂家货号已存在，默认跳过；更新请编辑原商品')
            codes.add(data['code']);factory_codes.add(pair)
        except (ValueError,TypeError) as e:error=str(e)
        result.append({'row':number,'data':data,'error':error})
    return {'headers':heads,'mapping':mapping,'rows':result,'warnings':warnings,'valid':sum(not r['error'] for r in result),'total':len(result)}

def literal(cell,value):
    cell.value=value
    if isinstance(value,str):cell.data_type='s'

def style_sheet(ws,header):
    for c in ws[header]:c.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='3566EC');c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[header].height=30
    ws.freeze_panes=f'A{header+1}'
    for col in range(1,ws.max_column+1):ws.column_dimensions[ws.cell(1,col).column_letter].width=19
    ws.column_dimensions['A'].width=28
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0

def template_bytes():
    w=Workbook();s=w.active;s.title='商品导入模板'
    s.append(list(FIELDS.values()));style_sheet(s,1)
    note=w.create_sheet('填写说明');notes=['在商品导入模板工作表填写数据，第一行是表头。','必填：销售编码（或导入时生成）、厂家编码、商品名称。其他字段可以留空。','销售编码和厂家编码建议设置为文本，保留前导零。','外箱尺寸请写明长、宽、高及单位；单箱重量单位为 kg，最多 3 位小数。产品参数可填写较长的文字。','未填写单件体积按 0 处理；提交前请核对。','出厂价单位：元/个；装箱数：个/件；体积：m³/件。','图片：主图和彩盒分别放在对应列、对应商品行；每组最多 5 张普通嵌入图片。复杂 WPS 图片公式可能需要手动补图。']
    for n in notes:note.append([n])
    note.column_dimensions['A'].width=100
    out=io.BytesIO();w.save(out);return out.getvalue()

def workbook_bytes(items,store,columns=None):
    from quotation import quotation_bytes
    return quotation_bytes(items,store,columns)


def export_bytes(items,store,mode,columns=None):
    if mode=='single':return workbook_bytes(items,store,columns),'.xlsx'
    if mode!='multiple':raise ValueError('导出方式无效')
    groups=defaultdict(list)
    for p in items:groups[p['factory']].append(p)
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        for i,(f,ps) in enumerate(groups.items(),1):
            name=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',f)[:70]
            z.writestr(f'{i}_{name}_报货单.xlsx',workbook_bytes(ps,store,columns))
    return out.getvalue(),'.zip'
