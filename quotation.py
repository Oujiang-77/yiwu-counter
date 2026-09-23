"""Configurable quotation columns and grouped image export."""
import io
import re
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU

EXPORT_FIELDS={'name':'商品名称','images':'主图','boxImages':'彩盒','code':'销售编码','factoryCode':'厂家编码','factory':'厂家名称','size':'尺寸','battery':'电池容量','charger':'充电头 / 接口','pack':'装箱数','quantity':'件数','units':'总数量','unitPrice':'单价','totalAmount':'总金额','volume':'单件体积','totalVolume':'总体积','cat':'分类','notes':'备注'}
EXPORT_FIELD_ALIASES={'price':'unitPrice','amount':'totalAmount'}

def validate_columns(columns):
    if not isinstance(columns,list) or not columns:raise ValueError('请至少选择一个有效的导出字段')
    columns=[EXPORT_FIELD_ALIASES.get(k,k) if isinstance(k,str) else k for k in columns]
    if any(not isinstance(k,str) or k not in EXPORT_FIELDS for k in columns):raise ValueError('请至少选择一个有效的导出字段')
    if len(columns)!=len(set(columns)):raise ValueError('导出字段不能重复')
    return columns

def quotation_bytes(items,store,columns=None):
    columns=validate_columns(list(EXPORT_FIELDS) if columns is None else columns)
    book=Workbook();book.remove(book.active);groups=defaultdict(list)
    for p in items:groups[p['factory']].append(p)
    for index,(factory,products) in enumerate(groups.items(),1):
        title=re.sub(r'[\\/*?:\[\]]','_',factory)[:25] or '厂家'
        if title in book.sheetnames:title+=f'_{index}'
        ws=book.create_sheet(title)
        ws.append(['厂家报价单',factory]);ws['B1'].data_type='s'
        ws.append(['说明','1 件 = 1 箱；总金额按单价计算，不含运费。'])
        ws.append([EXPORT_FIELDS[k] for k in columns])
        for c in ws[3]:
            c.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='3566EC')
        ws.row_dimensions[3].height=30
        for row,p in enumerate(products,4):
            height=32
            for col,key in enumerate(columns,1):
                if key in ('images','boxImages'):
                    names=p.get(key,([p['image']] if p.get('image') else []) if key=='images' else [])
                    height=max(height,58*len(names))
                    for slot,name in enumerate(names):
                        imagepath=store.root/'images'/name
                        if not imagepath.is_file():raise ValueError('商品图片缺失，请在资料中重新上传')
                        im=XLImage(str(imagepath));scale=min(100/im.width,68/im.height)
                        im.width*=scale;im.height*=scale
                        im.anchor=OneCellAnchor(_from=AnchorMarker(col=col-1,row=row-1,colOff=pixels_to_EMU(5),rowOff=pixels_to_EMU(slot*77+4)),ext=XDRPositiveSize2D(pixels_to_EMU(im.width),pixels_to_EMU(im.height)))
                        ws.add_image(im)
                    continue
                source_key={'unitPrice':'price','totalAmount':'amount'}.get(key,key)
                value=p.get(source_key,'')
                if key in ('totalAmount','totalVolume'):value=float(value)
                cell=ws.cell(row,col,value)
                if isinstance(value,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if key in ('unitPrice','totalAmount'):cell.number_format='0.00'
                if key in ('volume','totalVolume'):cell.number_format='0.000000'
            ws.row_dimensions[row].height=height
        last=3+len(products);total=last+1
        label_col=next((i for i,k in enumerate(columns,1) if k not in ('quantity','units','totalAmount','totalVolume')),None)
        if label_col:ws.cell(total,label_col,'合计')
        for col,key in enumerate(columns,1):
            letter=get_column_letter(col);ws.column_dimensions[letter].width=22 if key=='name' else 19
            if key in ('images','boxImages'):ws.column_dimensions[letter].width=17
            if key in ('quantity','units','totalAmount','totalVolume'):
                ws.cell(total,col,f'=SUM({letter}4:{letter}{last})')
                ws.cell(total,col).number_format='0.00' if key=='totalAmount' else '0.000000' if key=='totalVolume' else '0'
        ws.auto_filter.ref=f'A3:{get_column_letter(len(columns))}{last}'
        ws.freeze_panes='A4';ws.print_title_rows='1:3'
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    out=io.BytesIO();book.save(out);return out.getvalue()
